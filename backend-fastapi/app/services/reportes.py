"""Reporte diario de ventas en PDF y en Excel.

Los dos formatos salen exactamente de los mismos datos, los que devuelve
crud.ventas.reporte_diario, para que nunca puedan contar cosas distintas.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.factura import EMISOR

# La misma paleta del sitio, para que el documento se reconozca como de BIXE.
TINTA = colors.HexColor("#0d0f13")
MARCA = colors.HexColor("#0b7cc4")
SUAVE = colors.HexColor("#f4f6f8")
LINEA = colors.HexColor("#d9dde4")
GRIS = colors.HexColor("#5b626d")

MESES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)

ENCABEZADOS = (
    "N.º venta", "Hora", "Cliente", "Artículos", "Unidades", "Total", "Estado"
)


def fecha_larga(dia: date) -> str:
    return f"{dia.day} de {MESES[dia.month - 1]} de {dia.year}"


def _pesos(valor: float) -> str:
    """Formato colombiano: separador de miles con punto y sin decimales."""
    return f"$ {valor:,.0f}".replace(",", ".")


# ================================= PDF =================================


def generar_pdf(reporte: dict) -> bytes:
    memoria = BytesIO()
    documento = SimpleDocTemplate(
        memoria,
        pagesize=landscape(LETTER),
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Reporte de ventas {reporte['fecha']:%Y-%m-%d}",
        author=EMISOR["nombre"],
    )

    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "titulo", parent=estilos["Title"], fontSize=20, textColor=TINTA,
        alignment=0, spaceAfter=2,
    )
    sub = ParagraphStyle(
        "sub", parent=estilos["Normal"], fontSize=9.5, textColor=GRIS, leading=14
    )
    celda = ParagraphStyle(
        "celda", parent=estilos["Normal"], fontSize=8.5, leading=11
    )

    piezas = [
        Paragraph("Reporte diario de ventas", titulo),
        Paragraph(
            f"{EMISOR['nombre']} · {EMISOR['nit']}<br/>"
            f"{EMISOR['direccion']} · {EMISOR['ciudad']}",
            sub,
        ),
        Spacer(1, 8),
        Paragraph(
            f"<b>Fecha del reporte:</b> {fecha_larga(reporte['fecha'])} &nbsp;&nbsp;·&nbsp;&nbsp; "
            f"<b>Generado:</b> {reporte['generado_en']:%d/%m/%Y %H:%M}",
            sub,
        ),
        Spacer(1, 12),
    ]

    if not reporte["lineas"]:
        piezas.append(
            Paragraph("No se registraron ventas en esta fecha.", estilos["Normal"])
        )
        documento.build(piezas)
        return memoria.getvalue()

    filas = [list(ENCABEZADOS)]
    for linea in reporte["lineas"]:
        filas.append(
            [
                linea["numero"],
                linea["hora"],
                Paragraph(linea["cliente"], celda),
                Paragraph(linea["articulos"], celda),
                str(linea["unidades"]),
                _pesos(linea["total"]),
                linea["estado"].capitalize(),
            ]
        )

    tabla = Table(
        filas,
        colWidths=[26 * mm, 15 * mm, 48 * mm, 80 * mm, 20 * mm, 30 * mm, 24 * mm],
        repeatRows=1,
    )
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), TINTA),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (4, 0), (5, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, SUAVE]),
                ("GRID", (0, 0), (-1, -1), 0.4, LINEA),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    piezas += [tabla, Spacer(1, 14)]

    resumen = Table(
        [
            ["Ventas registradas", str(reporte["total_ventas"])],
            ["Unidades vendidas", str(reporte["unidades"])],
            ["Base gravable", _pesos(reporte["subtotal"])],
            ["Descuentos", _pesos(reporte["descuento"])],
            ["IVA (19%)", _pesos(reporte["impuesto"])],
            ["Total del día", _pesos(reporte["total"])],
        ],
        colWidths=[45 * mm, 38 * mm],
        hAlign="RIGHT",
    )
    resumen.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("TEXTCOLOR", (0, 0), (0, -1), GRIS),
                ("LINEABOVE", (0, -1), (-1, -1), 0.8, TINTA),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, -1), (-1, -1), MARCA),
                ("FONTSIZE", (0, -1), (-1, -1), 11),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    piezas.append(resumen)

    piezas += [
        Spacer(1, 16),
        Paragraph(
            "Las ventas anuladas aparecen en el listado pero no suman en los "
            "totales. Documento generado automáticamente por el sistema BIXE.",
            ParagraphStyle("pie", parent=sub, fontSize=7.5, textColor=GRIS),
        ),
    ]

    documento.build(piezas)
    return memoria.getvalue()


# ================================ Excel ================================


def generar_excel(reporte: dict) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Ventas"

    negrita_blanca = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="0D0F13")
    borde = Border(*(Side(style="thin", color="D9DDE4"),) * 4)
    pesos = '"$" #,##0'

    hoja["A1"] = "Reporte diario de ventas"
    hoja["A1"].font = Font(bold=True, size=15)
    hoja["A2"] = f"{EMISOR['nombre']} · {EMISOR['nit']}"
    hoja["A3"] = f"Fecha del reporte: {fecha_larga(reporte['fecha'])}"
    hoja["A4"] = f"Generado: {reporte['generado_en']:%d/%m/%Y %H:%M}"
    for fila in (2, 3, 4):
        hoja[f"A{fila}"].font = Font(size=9, color="5B626D")

    # Fila 6 es el encabezado de la tabla; la 7 es donde arrancan los datos.
    for columna, titulo in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=6, column=columna, value=titulo)
        celda.font = negrita_blanca
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde

    for indice, linea in enumerate(reporte["lineas"], start=7):
        valores = (
            linea["numero"],
            linea["hora"],
            linea["cliente"],
            linea["articulos"],
            linea["unidades"],
            linea["total"],
            linea["estado"].capitalize(),
        )
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=indice, column=columna, value=valor)
            celda.border = borde
            if columna == 6:
                celda.number_format = pesos
            if columna == 4:
                celda.alignment = Alignment(wrap_text=True, vertical="top")

    ultima = 6 + len(reporte["lineas"])

    # El autofiltro es lo que permite ordenar y filtrar después en Excel, que
    # es justo para lo que pide el requerimiento que exista la exportación.
    hoja.auto_filter.ref = f"A6:{get_column_letter(len(ENCABEZADOS))}{max(ultima, 7)}"
    hoja.freeze_panes = "A7"

    for columna, ancho in enumerate((14, 8, 28, 52, 10, 16, 12), start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    fila = ultima + 2
    for etiqueta, valor, formato in (
        ("Ventas registradas", reporte["total_ventas"], None),
        ("Unidades vendidas", reporte["unidades"], None),
        ("Base gravable", reporte["subtotal"], pesos),
        ("Descuentos", reporte["descuento"], pesos),
        ("IVA (19%)", reporte["impuesto"], pesos),
        ("Total del día", reporte["total"], pesos),
    ):
        hoja.cell(row=fila, column=5, value=etiqueta).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=6, value=valor)
        if formato:
            celda.number_format = formato
        fila += 1

    hoja.cell(row=fila - 1, column=5).font = Font(bold=True, size=12)
    hoja.cell(row=fila - 1, column=6).font = Font(bold=True, size=12, color="0B7CC4")

    memoria = BytesIO()
    libro.save(memoria)
    return memoria.getvalue()
